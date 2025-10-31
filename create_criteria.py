from docx import Document

# 创建评分标准文档
doc = Document()
doc.add_heading('应用文评分标准', 0)

doc.add_heading('评分维度（总分100分）', level=1)

doc.add_heading('1. 内容完整性（25分）', level=2)
doc.add_paragraph('是否包含所有要求的要点')
doc.add_paragraph('要点是否清晰、准确')
doc.add_paragraph('内容是否充实、具体')

doc.add_heading('2. 语法准确性（25分）', level=2)
doc.add_paragraph('时态、语态使用是否正确')
doc.add_paragraph('句子结构是否完整、准确')
doc.add_paragraph('标点符号使用是否规范')

doc.add_heading('3. 词汇使用（25分）', level=2)
doc.add_paragraph('词汇选择是否准确、得体')
doc.add_paragraph('词汇使用是否丰富、多样')
doc.add_paragraph('拼写是否正确')

doc.add_heading('4. 组织结构（15分）', level=2)
doc.add_paragraph('文章结构是否清晰、合理')
doc.add_paragraph('段落之间是否连贯')
doc.add_paragraph('逻辑是否清晰')

doc.add_heading('5. 格式规范（10分）', level=2)
doc.add_paragraph('是否符合应用文格式要求')
doc.add_paragraph('称呼、落款是否规范')
doc.add_paragraph('整体排版是否美观')

doc.add_heading('评分等级', level=1)
doc.add_paragraph('优秀（90-100分）：各项指标均达到优秀水平')
doc.add_paragraph('良好（75-89分）：各项指标基本达标，有少量小问题')
doc.add_paragraph('中等（60-74分）：部分指标有明显问题，需要改进')
doc.add_paragraph('不及格（0-59分）：多项指标存在严重问题')

doc.save('应用文评分标准.docx')
print('评分标准文档创建成功！')
